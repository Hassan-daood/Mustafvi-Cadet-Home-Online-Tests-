"""Convert the supplied PAF MCQ document into this app's normal-import format."""
import re
from pathlib import Path

from docx import Document


SOURCE = Path(r"C:\Users\Ahmad Raza\Downloads\PAF_Sargodha_Lower_Topa_100_Competition_MCQs_All_Topics.docx")
OUTPUT = Path(__file__).resolve().parent / "PAF_Sargodha_Lower_Topa_normal_import.docx"

# Answer key checked against the four choices in the supplied 100-question paper.
ANSWERS = (
    "BBBBB" "BBBBB" "BABBB" "AABAB" "BACDA"
    "BBBBC" "BBCBB" "BBBBA" "BBBAB" "BABBA"
    "BBBBB" "AABAB" "BBBBB" "BBABA" "BABAB"
    "BCBAB" "CBACA" "BAABC" "BAAAA" "BABAB"
)


def option_values(line):
    matches = re.findall(r"(?i)([abcd])\)\s*(.*?)(?=\s+[abcd]\)\s|$)", line)
    values = {letter.upper(): text.strip() for letter, text in matches}
    return values if set(values) == {"A", "B", "C", "D"} else None


def main():
    doc = Document(SOURCE)
    lines = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
    extracted = []
    section = "Imported Section"
    index = 0

    while index < len(lines):
        line = lines[index]
        question = re.match(r"^(\d+)\.\s+(.+)$", line)
        options = option_values(lines[index + 1]) if index + 1 < len(lines) else None
        if question and options:
            extracted.append({"section": section, "question": question.group(2), "options": options})
            index += 2
        elif re.match(r"^\d+\.\s+", line):
            section = re.sub(r"^\d+\.\s+", "", line)
            index += 1
        else:
            index += 1

    if len(extracted) != 100 or len(ANSWERS) != 100:
        raise RuntimeError(f"Expected 100 questions and answers; found {len(extracted)} questions and {len(ANSWERS)} answers.")

    output = Document()
    output.add_heading("PAF Sargodha / Lower Topa - Normal Import", level=1)
    current_section = None
    for item, answer in zip(extracted, ANSWERS):
        if item["section"] != current_section:
            current_section = item["section"]
            output.add_paragraph(f"SECTION: {current_section}")
            output.add_paragraph("TIME: 10")
        output.add_paragraph(f"Q: {item['question']}")
        for letter in "ABCD":
            output.add_paragraph(f"{letter}: {item['options'][letter]}")
        output.add_paragraph(f"ANSWER: {answer}")

    output.save(OUTPUT)
    print(f"Created: {OUTPUT}")
    print(f"Questions: {len(extracted)} | Sections: {len(set(q['section'] for q in extracted))}")


if __name__ == "__main__":
    main()
